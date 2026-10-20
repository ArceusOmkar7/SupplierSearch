"""Import actual official exports and look up exact identifiers.

Only legitimately obtained files with a validated manifest are imported; original bytes and
the source cutoff are preserved, rows link to their page/row, and an import today never
refreshes the status the export described. This is a known-supplier corpus, not discovery.
"""

import csv
import io
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

from bs4 import BeautifulSoup
from sqlalchemy import select
from sqlalchemy.orm import Session

from verisource.db.models import SourceDataset, SourceDatasetRecord
from verisource.domain.intake import json_hash
from verisource.sources.manifests import load_record
from verisource.storage.filesystem import PrivateFileStore

SCHEMA_VERSION = "tabular-v1"
# Header text (normalized) -> identifier type used for exact lookup.
IDENTIFIER_COLUMNS = {
    "licence no": "licence_number",
    "license no": "licence_number",
    "gstin": "gstin",
    "cin": "cin",
    "llpin": "llpin",
    "udyam": "udyam",
    "udyam registration number": "udyam",
}
NAME_COLUMNS = {"firm name", "company name", "name of the company", "legal name"}
# Direct personal contact columns are never copied into case views or free inference.
EXCLUDED_COLUMNS = {"contact details", "email", "phone", "mobile", "contact"}


@dataclass(frozen=True)
class ImportResult:
    dataset_id: UUID
    imported: int
    quarantined: int
    issues: list[str]


@dataclass(frozen=True)
class LookupMatch:
    record_id: UUID
    dataset_id: UUID
    locator: str
    fields: dict[str, Any]
    publication_cutoff: datetime | None
    obtained_at: datetime
    historical: bool
    current_status_established: bool = False


def _clean(value: Any) -> str:
    return " ".join(str(value or "").split())


def _header_key(value: Any) -> str:
    return _clean(value).casefold().replace("_", " ")


def _rows_from_pdf(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    import pdfplumber

    tables: list[tuple[str, list[list[Any]]]] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for index, page in enumerate(pdf.pages, start=1):
            for table in page.extract_tables():
                tables.append((f"page {index}", table))
    return tables


def _rows_from_csv(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    text = content.decode("utf-8-sig")
    return [("sheet 1", [list(row) for row in csv.reader(io.StringIO(text))])]


def _rows_from_xlsx(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    return [
        (f"sheet {sheet.title}", [list(row) for row in sheet.iter_rows(values_only=True)])
        for sheet in workbook.worksheets
    ]


def _rows_from_html(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    soup = BeautifulSoup(content, "html.parser")
    result: list[tuple[str, list[list[Any]]]] = []
    for number, table in enumerate(soup.find_all("table"), start=1):
        rows = [
            [cell.get_text(" ") for cell in row.find_all(["th", "td"])]
            for row in table.find_all("tr")
        ]
        result.append((f"table {number}", rows))
    return result


def import_source_dataset(
    db: Session,
    store: PrivateFileStore,
    manifest_path: Path,
    *,
    records_root: Path | None = None,
) -> ImportResult:
    manifest, content = load_record(manifest_path, records_root)
    fmt = manifest["format"]
    if fmt not in {"pdf", "csv", "xlsx", "html"}:
        raise ValueError("Unsupported dataset format")
    existing = db.scalar(
        select(SourceDataset).where(
            SourceDataset.sha256 == manifest["sha256"],
            SourceDataset.schema_version == SCHEMA_VERSION,
        )
    )
    if existing is not None:
        return ImportResult(
            existing.id, existing.imported_count, existing.quarantined_count, ["already imported"]
        )
    reader = {
        "pdf": _rows_from_pdf,
        "csv": _rows_from_csv,
        "xlsx": _rows_from_xlsx,
        "html": _rows_from_html,
    }[fmt]
    dataset_id = uuid4()
    stored = store.put(dataset_id, content)
    records: list[SourceDatasetRecord] = []
    quarantined = 0
    issues: list[str] = []
    for location, table in reader(content):
        header_index = None
        for index, row in enumerate(table):
            keys = {_header_key(cell) for cell in row}
            if keys & set(IDENTIFIER_COLUMNS):
                header_index = index
                break
        if header_index is None:
            issues.append(f"{location}: no recognizable identifier header; table quarantined")
            quarantined += max(len(table) - 1, 0)
            continue
        header = [_header_key(cell) for cell in table[header_index]]
        for offset, row in enumerate(table[header_index + 1 :], start=header_index + 2):
            if not any(_clean(cell) for cell in row):
                continue
            identifiers: dict[str, str] = {}
            fields: dict[str, Any] = {}
            for key, cell in zip(header, row, strict=False):
                value = _clean(cell)
                if not value or key in EXCLUDED_COLUMNS:
                    continue
                if key in IDENTIFIER_COLUMNS:
                    identifiers[IDENTIFIER_COLUMNS[key]] = value.upper()
                if key in NAME_COLUMNS:
                    fields["firm_name"] = value
                elif key in {"factory address", "address", "registered office address"}:
                    fields["factory_address"] = value
                elif key == "validity":
                    fields["validity"] = value
                elif key not in IDENTIFIER_COLUMNS:
                    fields[key.replace(" ", "_")] = value
            if not identifiers or "firm_name" not in fields:
                quarantined += 1
                issues.append(f"{location} row {offset}: missing identifier or name")
                continue
            records.append(
                SourceDatasetRecord(
                    dataset_id=dataset_id,
                    locator=f"{location}, row {offset}",
                    normalized_identifiers=identifiers,
                    selected_business_fields=fields,
                    origin_hash=json_hash([identifiers, fields]),
                )
            )
    obtained = datetime.fromisoformat(manifest["obtained_at"])
    cutoff = (
        datetime.fromisoformat(manifest["publication_cutoff"])
        if manifest["publication_cutoff"]
        else None
    )
    db.add(
        SourceDataset(
            id=dataset_id,
            source_version=f"{manifest['publisher']}@{manifest['obtained_at']}",
            publisher=manifest["publisher"],
            original_url=manifest["original_url"],
            licence_attribution=f"{manifest['rights']} | {manifest['attribution']}",
            obtained_at=obtained,
            publication_cutoff=cutoff,
            cutoff_unknown_reason=manifest["cutoff_unknown_reason"],
            coverage=manifest["coverage"],
            local_storage_key=stored.storage_key,
            sha256=stored.sha256,
            format=fmt,
            schema_version=SCHEMA_VERSION,
            state="imported" if records else "quarantined",
            imported_count=len(records),
            quarantined_count=quarantined,
            imported_at=datetime.now(UTC),
        )
    )
    db.flush()
    db.add_all(records)
    db.flush()
    return ImportResult(dataset_id, len(records), quarantined, issues)


def lookup_exact(
    db: Session,
    *,
    identifier_type: str,
    value: str,
    dataset_id: UUID | None = None,
) -> list[LookupMatch]:
    """Exact normalized identifier match only; no fuzzy or prefix matching, no discovery."""
    query = (
        select(SourceDatasetRecord, SourceDataset)
        .join(SourceDataset, SourceDataset.id == SourceDatasetRecord.dataset_id)
        .where(
            SourceDataset.state == "imported",
            SourceDatasetRecord.normalized_identifiers.contains(
                {identifier_type: value.strip().upper()}
            ),
        )
        .order_by(SourceDataset.obtained_at.desc(), SourceDatasetRecord.locator)
    )
    if dataset_id is not None:
        query = query.where(SourceDataset.id == dataset_id)
    return [
        LookupMatch(
            record_id=record.id,
            dataset_id=dataset.id,
            locator=record.locator,
            fields=dict(record.selected_business_fields),
            publication_cutoff=dataset.publication_cutoff,
            obtained_at=dataset.obtained_at,
            historical=dataset.publication_cutoff is not None
            and dataset.publication_cutoff.date() < dataset.obtained_at.date(),
        )
        for record, dataset in db.execute(query)
    ]
